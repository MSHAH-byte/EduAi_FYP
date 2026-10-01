import json
import tempfile
from pathlib import Path
import httpx
from docx import Document as DocxDocument

base = 'http://127.0.0.1:8000'

# OpenAPI route presence
openapi = httpx.get(f'{base}/openapi.json', timeout=20).json()
print('OPENAPI_PATHS', sorted(openapi['paths'].keys()))

# Document success case via DOCX
with tempfile.TemporaryDirectory() as tmp:
    docx_path = Path(tmp) / 'sample.docx'
    doc = DocxDocument()
    doc.add_heading('Course Overview', level=1)
    doc.add_paragraph('This document contains a paragraph and a table.')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Topic'
    table.cell(0, 1).text = 'Value'
    doc.save(docx_path)
    with docx_path.open('rb') as f:
        files = {'file': ('sample.docx', f, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}
        r = httpx.post(f'{base}/api/v1/document/summarize', files=files, timeout=240)
    print('DOCUMENT_STATUS', r.status_code)
    print('DOCUMENT_BODY', r.text)

# Existing chat endpoint
chat_r = httpx.post(f'{base}/api/v1/chat/', json={'message':'hello','conversation_history':[]}, timeout=240)
print('CHAT_STATUS', chat_r.status_code)
print('CHAT_BODY', chat_r.text)

# Existing generate notes endpoint
notes_r = httpx.post(f'{base}/api/v1/generate/notes', json={'topic':'photosynthesis','num_items':5}, timeout=240)
print('NOTES_STATUS', notes_r.status_code)
print('NOTES_BODY', notes_r.text)
